"""The AdrianMath brand header — one look per series (Adrian, 17 Sep 2026).

"keep the design B for A Math and design A for E Math - but make it more obvious
which is A Math and which is E Math … For the footer remove Adrian Fong … are you
able to come up with design for Sec 1 and Sec 2? same idea, slightly different
designs for each of them (so it's easy to distinguish the worksheets)"

Every series carries the same pieces: the √m mark + AdrianMath / TUITION lock-up,
the level line, a coloured SUBJECT BLOCK at the right end of the header ("A MATH",
"E MATH", "SEC 1", "SEC 2") that reads from across the room, the topic title, a
PRACTICE · n questions · m marks line, Name/Date, a running header from page 2 that
repeats the subject in its colour, and the footer "AdrianMath Tuition ·
adrianmathtuition.com | Page x of y". What changes per series is the header
treatment and the accent colour:

    series  header treatment                       accent
    AM      navy band across the page (design B)   orange
    EM      white, rule under the header (design A) teal
    S1      pale green tinted band                  green
    S2      white, thick blue bar over the header   bright blue

Every series also has a BLACK-AND-WHITE version (`mono=True`, Adrian: "i usually
print in black and white, can i have other versions without colour? that means
the design itself will distinguish the papers"). There the structure alone tells
the sheets apart — no colour is relied on:

    series  header                                   subject block          page-2 rule
    AM      white, thick rule over + thin rule under small solid black tab  thick solid
    EM      white, one rule under the header         outlined box           double
    S1      near-white grey tint (F2F2F2)            double-lined box       dotted
    S2      white, thick bar over the header         heavy rules above and  dashed
                                                     below, sides open

(The first mono set, a solid black A Math band, was dropped the same day —
Adrian: "too oppressive as black, and waste ink when printing".) The mark in
black and white is the outlined ring (`mark_outline.png`), not a filled disc.

Use it through `Worksheet.brand(level, topic, ...)`. It inserts at the TOP of the
body, so it can be called before or after the questions are written (after is
handier: the question count and marks are known then). A branded sheet does not
also call `ws.title()` / `ws.subtitle()`.
"""
from pathlib import Path

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ASSETS = Path(__file__).parent / 'assets' / 'brand'

NAVY = '1B2A4A'
ORANGE = 'E08A3A'   # the "Math" in AdrianMath — brand, on every series
GREY = '6E7A8E'
PALE = 'C8D0DE'     # small print on the navy band
RULE = 'BEC6D2'
WHITE = 'FFFFFF'
TEXT_W_CM = 16.0    # A4 less worksheet_lib's 2.5 cm margins

# The palette every piece reads. `ground` = the header band's fill (None = white),
# `block` = the subject block's fill (None = an outlined box), `frame` = the
# masthead's own rule, `run_rule` = the rule under the page-2 running header.
COLOUR = dict(ink=NAVY, math=ORANGE, grey=GREY, pale=PALE, rule=RULE,
              band_math=ORANGE, mark='mark_navy.png', band_mark='mark_white.png')
MONO = dict(ink='1A1A1A', math='808080', grey='666666', pale='BFBFBF', rule='BFBFBF',
            band_math='A6A6A6', mark='mark_outline.png', band_mark='mark_white.png')

SERIES = {
    'AM': dict(tag='A MATH', style='band', ground=NAVY,
               accent=ORANGE, block=ORANGE, tag_ink=NAVY, title_bar=ORANGE,
               run_rule=('single', 8, ORANGE)),
    'EM': dict(tag='E MATH', style='plain', frame=('bottom', 'single', 18, '0E8A7D'),
               accent='0E8A7D', block='0E8A7D', tag_ink=WHITE,
               run_rule=('single', 8, '0E8A7D')),
    'S1': dict(tag='SEC 1', style='tint', ground='E9F5ED',
               accent='2E9A58', block='2E9A58', tag_ink=WHITE, title_bar='2E9A58',
               run_rule=('single', 8, '2E9A58')),
    # Sec 2 was plum; Adrian 17 Sep 2026: "change purple, something suitable for
    # secondary school students"
    'S2': dict(tag='SEC 2', style='bar', frame=('top', 'single', 36, '1F74D6'),
               accent='1F74D6', block='1F74D6', tag_ink=WHITE,
               run_rule=('single', 8, '1F74D6')),
}

# The black-and-white set. Adrian 17 Sep 2026, after the first mono set (a solid
# black A Math band): "the large black background is too oppressive as black, and
# waste ink when printing". So no band is filled: every series is white or a
# near-white tint, the mark is the outlined ring, and the SHAPE of the rules and
# of the subject block is what tells the sheets apart. `frame` may be a list of
# rules; `block_rules` draws rules round the subject block only (None = no rule).
SERIES_MONO = {
    # ruled top and bottom like a newspaper masthead; the one small solid tab
    'AM': dict(tag='A MATH', style='plain',
               frame=[('top', 'single', 24, '1A1A1A'), ('bottom', 'single', 6, '1A1A1A')],
               accent='1A1A1A', block='1A1A1A', tag_ink=WHITE, title_bar='1A1A1A',
               run_rule=('single', 12, '1A1A1A')),
    # one rule under the header; the subject in an outlined box
    'EM': dict(tag='E MATH', style='plain', frame=('bottom', 'single', 12, '1A1A1A'),
               accent='1A1A1A', block=None, tag_ink='1A1A1A',
               block_rules={s: ('single', 12) for s in ('top', 'left', 'bottom', 'right')},
               run_rule=('double', 6, '1A1A1A')),
    # a whisper of grey behind the header; the subject in a double box
    'S1': dict(tag='SEC 1', style='tint', ground='F2F2F2',
               accent='1A1A1A', block=WHITE, tag_ink='1A1A1A', title_bar='A6A6A6',
               block_rules={s: ('double', 6) for s in ('top', 'left', 'bottom', 'right')},
               run_rule=('dotted', 12, '1A1A1A')),
    # a thick bar over the header; the subject between two heavy rules, sides open
    'S2': dict(tag='SEC 2', style='bar', frame=('top', 'single', 36, '1A1A1A'),
               accent='1A1A1A', block=None, tag_ink='1A1A1A',
               block_rules={'top': ('single', 18), 'bottom': ('single', 18)},
               run_rule=('dashed', 12, '1A1A1A')),
}

# questions.level -> (series, the small line under the subject block, level line)
LEVELS = {
    'AM':       ('AM', 'SEC 4', 'Sec 4 Additional Mathematics'),
    'S3_AM':    ('AM', 'SEC 3', 'Sec 3 Additional Mathematics'),
    'AM_NA':    ('AM', 'SEC 5 N(A)', 'Sec 5 N(A) Additional Mathematics'),
    'EM':       ('EM', 'SEC 4', 'Sec 4 Mathematics'),
    'S3_EM':    ('EM', 'SEC 3', 'Sec 3 Mathematics'),
    'EM_NA':    ('EM', 'SEC 4 N(A)', 'Sec 4 N(A) Mathematics'),
    'S3_EM_NA': ('EM', 'SEC 3 N(A)', 'Sec 3 N(A) Mathematics'),
    'S3_EM_NT': ('EM', 'SEC 3 N(T)', 'Sec 3 N(T) Mathematics'),
    'S1':       ('S1', 'MATHEMATICS', 'Sec 1 Mathematics'),
    'S2':       ('S2', 'MATHEMATICS', 'Sec 2 Mathematics'),
}


def _rgb(hex_):
    return RGBColor.from_string(hex_)


def _run(p, text, size=9, bold=False, color=NAVY, font='Arial', spacing=None):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = _rgb(color)
    rPr = r._element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    for a in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        rf.set(qn(f'w:{a}'), font)
    if spacing:
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:val'), str(spacing))
        rPr.append(sp)
    return r


def _tight(p, before=0, after=0, align=None):
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = 1.0
    if align:
        p.alignment = align
    return p


def _para_border(p, side, color, sz, space, val='single'):
    pPr = p._p.get_or_add_pPr()
    b = pPr.find(qn('w:pBdr'))
    if b is None:
        b = OxmlElement('w:pBdr')
        pPr.append(b)
    e = OxmlElement(f'w:{side}')
    for k, v in (('val', val), ('sz', sz), ('space', space), ('color', color)):
        e.set(qn(f'w:{k}'), str(v))
    b.append(e)


def _table_borders(tbl, **sides):
    """sides: name -> (sz, color) or (val, sz, color); every other side is nil."""
    tblPr = tbl._tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{s}')
        if s in sides:
            val, sz, color = sides[s] if len(sides[s]) == 3 else ('single', *sides[s])
            for k, v in (('val', val), ('sz', sz), ('space', 0), ('color', color)):
                e.set(qn(f'w:{k}'), str(v))
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)
    tblPr.append(b)


def _cell_margins(tbl, top=0, bottom=0, left=0, right=0):
    m = OxmlElement('w:tblCellMar')
    for s, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{s}')
        e.set(qn('w:w'), str(int(v)))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tbl._tbl.tblPr.append(m)


def _widths(tbl, cms):
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    for gc, w in zip(tbl._tbl.tblGrid.findall(qn('w:gridCol')), cms):
        gc.set(qn('w:w'), str(int(Cm(w).twips)))
    for row in tbl.rows:
        for c, w in zip(row.cells, cms):
            c.width = Cm(w)
    old = tblPr.find(qn('w:tblW'))
    if old is not None:
        tblPr.remove(old)
    tw = OxmlElement('w:tblW')
    tw.set(qn('w:w'), str(int(Cm(sum(cms)).twips)))
    tw.set(qn('w:type'), 'dxa')
    tblPr.append(tw)


def _shade(cell, fill):
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(s)


def _cell_box(cell, rules, color):
    """Rules round one cell (the mono subject blocks): side -> (val, sz)."""
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for s in ('top', 'left', 'bottom', 'right'):
        if s not in rules:
            continue
        e = OxmlElement(f'w:{s}')
        val, sz = rules[s]
        for k, v in (('val', val), ('sz', sz), ('space', 0), ('color', color)):
            e.set(qn(f'w:{k}'), str(v))
        b.append(e)
    tcPr.append(b)


def _field(p, code, **kw):
    """PAGE / NUMPAGES as a real Word field."""
    def char(kind, text=''):
        r = _run(p, text, **kw)
        c = OxmlElement('w:fldChar')
        c.set(qn('w:fldCharType'), kind)
        r._element.insert(1, c)
    char('begin')
    r = _run(p, '', **kw)
    t = OxmlElement('w:instrText')
    t.set(qn('xml:space'), 'preserve')
    t.text = f' {code} '
    r._element.append(t)
    char('separate')
    _run(p, '1', **kw)
    char('end')


# ── the header lock-up ─────────────────────────────────────────────────────

def _masthead(doc, cfg, k, small, level_line):
    band = cfg['style'] == 'band'
    ground = cfg.get('ground')
    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    frame = cfg.get('frame') or []
    if isinstance(frame, tuple):
        frame = [frame]
    _table_borders(tbl, **{side: (val, sz, color) for side, val, sz, color in frame})
    pad = 170 if ground else 110
    _cell_margins(tbl, top=pad, bottom=pad, left=160 if ground else 0, right=160 if ground else 0)
    _widths(tbl, (2.0, 4.5, 6.1, 3.4))
    logo, word, level, block = tbl.rows[0].cells
    for c in tbl.rows[0].cells:
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ground:
            _shade(c, ground)
    if cfg['block']:
        _shade(block, cfg['block'])
    if cfg.get('block_rules'):
        _cell_box(block, cfg['block_rules'], cfg['tag_ink'])

    p = _tight(logo.paragraphs[0])
    mark = k['band_mark'] if band else k['mark']
    p.add_run().add_picture(str(ASSETS / mark), width=Cm(1.35))

    p = _tight(word.paragraphs[0])
    _run(p, 'Adrian', size=17, bold=True, color=WHITE if band else k['ink'])
    _run(p, 'Math', size=17, bold=True, color=k['band_math'] if band else k['math'])
    p = _tight(word.add_paragraph())
    _run(p, 'TUITION', size=7.5, bold=True, color=k['pale'] if band else k['grey'], spacing=60)

    # keep the level text off the subject block's edge
    p = _tight(level.paragraphs[0], align=WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.right_indent = Cm(0.35)
    _run(p, level_line, size=10 if len(level_line) <= 22 else 9, bold=True,
         color=WHITE if band else k['ink'])
    p = _tight(level.add_paragraph(), align=WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.right_indent = Cm(0.35)
    _run(p, 'adrianmathtuition.com', size=8, color=k['pale'] if band else k['grey'])

    # the subject block: the one thing that says which sheet this is
    p = _tight(block.paragraphs[0], align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, cfg['tag'], size=17, bold=True, color=cfg['tag_ink'], spacing=10)
    p = _tight(block.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, small, size=7, bold=True, color=cfg['tag_ink'], spacing=40)
    return tbl._tbl


def _title_block(doc, cfg, k, topic, kind, n, marks):
    accent = cfg['accent']
    bar = cfg.get('title_bar')
    els = []
    p = _tight(doc.add_paragraph(), before=12, after=1)
    if bar:
        _para_border(p, 'left', bar, 36, 8)
    _run(p, topic, font='Georgia', size=19, bold=True, color=k['ink'])
    els.append(p._p)
    p = _tight(doc.add_paragraph(), after=4)
    if bar:
        _para_border(p, 'left', bar, 36, 8)
    _run(p, kind.upper(), size=9, bold=True, color=accent, spacing=20)
    bits = []
    if n:
        bits.append(f'{n} question' + ('' if n == 1 else 's'))
    if marks:
        bits.append(f'{marks} marks')
    if bits:
        _run(p, '   ·   ' + '   ·   '.join(bits), size=9, color=k['grey'])
    els.append(p._p)
    p = _tight(doc.add_paragraph(), after=8)
    _run(p, 'Name ', size=8.5, color=k['grey'])
    _run(p, '_' * 38, size=8.5, color=k['rule'])
    _run(p, '      Date ', size=8.5, color=k['grey'])
    _run(p, '_' * 16, size=8.5, color=k['rule'])
    _para_border(p, 'bottom', k['rule'], 4, 6)
    els.append(p._p)
    return els


# ── page furniture ─────────────────────────────────────────────────────────

def _two_sided(part, left, right, side, color, sz, val='single'):
    """A left/right line in a header or footer: a borderless two-cell table, the
    part's own empty paragraph kept after it at 1 pt."""
    tbl = part.add_table(rows=1, cols=2, width=Cm(TEXT_W_CM))
    _table_borders(tbl, **{side: (val, sz, color)})
    _cell_margins(tbl, top=50 if side == 'top' else 0, bottom=50 if side == 'bottom' else 0)
    _widths(tbl, (10.5, 5.5))
    lc, rc = tbl.rows[0].cells
    left(_tight(lc.paragraphs[0]))
    right(_tight(rc.paragraphs[0], align=WD_ALIGN_PARAGRAPH.RIGHT))
    first = part.paragraphs[0]
    part._element.remove(tbl._tbl)
    first._p.addprevious(tbl._tbl)
    _tight(first)
    first.paragraph_format.line_spacing = Pt(1)
    first.add_run('').font.size = Pt(1)


def _furniture(doc, cfg, k, small, topic):
    s = doc.sections[0]
    s.bottom_margin = Cm(1.7)
    s.footer_distance = Cm(0.7)
    s.header_distance = Cm(0.9)
    s.different_first_page_header_footer = True

    def site(p):
        _run(p, 'AdrianMath Tuition', size=8, bold=True, color=k['ink'])
        _run(p, '  ·  adrianmathtuition.com', size=8, color=k['grey'])

    def page(p):
        _run(p, 'Page ', size=8, color=k['grey'])
        _field(p, 'PAGE', size=8, color=k['grey'])
        _run(p, ' of ', size=8, color=k['grey'])
        _field(p, 'NUMPAGES', size=8, color=k['grey'])

    for ft in (s.first_page_footer, s.footer):
        _two_sided(ft, site, page, 'top', k['rule'], 4)

    def brand(p):
        _run(p, 'Adrian', size=8, bold=True, color=k['ink'])
        _run(p, 'Math', size=8, bold=True, color=k['math'])
        _run(p, '  ·  ' + topic, size=8, color=k['grey'])

    def subject(p):
        _run(p, cfg['tag'], size=8.5, bold=True, color=cfg['accent'], spacing=10)
        if small != 'MATHEMATICS':
            _run(p, '  ·  ' + small.title().replace('N(a)', 'N(A)').replace('N(t)', 'N(T)'),
                 size=8, color=k['grey'])

    val, sz, color = cfg['run_rule']
    _two_sided(s.header, brand, subject, 'bottom', color, sz, val)


def apply(ws, level, topic, kind='Practice', n_questions=None, marks=None, mono=False):
    """Brand a Worksheet for `level` (a questions.level value, e.g. 'AM', 'S3_EM',
    'S1'). Inserts the masthead + title block at the top of the body and sets
    the header/footer. `mono=True` = the black-and-white version of the series.
    Returns the series key."""
    if level not in LEVELS:
        raise ValueError(f'no brand design for level {level!r} — known: {", ".join(LEVELS)}')
    series, small, level_line = LEVELS[level]
    cfg = (SERIES_MONO if mono else SERIES)[series]
    k = MONO if mono else COLOUR
    doc = ws.doc
    els = [_masthead(doc, cfg, k, small, level_line)]
    els += _title_block(doc, cfg, k, topic, kind, n_questions, marks)
    body = doc.element.body
    for i, el in enumerate(els):
        body.remove(el)
        body.insert(i, el)
    _furniture(doc, cfg, k, small, topic)
    if mono:
        _drain_colour(doc)
    return series


def _drain_colour(doc):
    """The black-and-white version carries no colour anywhere: the orange [Ans:]
    lines and any coloured working in the body and the styles go dark grey."""
    for root in (doc.element.body, doc.styles.element):
        for c in root.iter(qn('w:color')):
            v = (c.get(qn('w:val')) or '').upper()
            if len(v) == 6 and not (v[0:2] == v[2:4] == v[4:6]):
                c.set(qn('w:val'), '404040')
                for t in ('w:themeColor', 'w:themeShade', 'w:themeTint'):
                    if c.get(qn(t)) is not None:
                        del c.attrib[qn(t)]
