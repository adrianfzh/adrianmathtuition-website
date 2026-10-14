"""Solutions as a TABLE — Adrian's solutions layout (ADRIAN-STYLE.md §10, 14 Oct 2026).

    python3 solutions_table.py <build_script.py> <out.docx> ["Title above the table"]

<build_script.py> is an ordinary worksheet build script written against
worksheet_lib's Worksheet API (ws.Q(...) then ws.solution_box([(label, steps), ...])).
This file RECORDS those calls instead of rendering a boxed sheet, then lays every
question out as one 3-column table: number | part | working, one row per part,
one step per line (Word equations, "=" continuation lines indented under the "="),
the solution's "Answer: …" row becomes a bold "∴ …" conclusion line.
First used for "A REV 2026 5 Vectors (Practice Solutions) (table).docx".
Run from this folder (it imports worksheet_lib) with python-docx + lxml installed.
"""
import re, sys, copy, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import worksheet_lib as wl
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

QUESTIONS = []   # [(rows)]
class Recorder:
    def __init__(self, *a, **k): pass
    def title(self, t): pass
    def subtitle(self, t): pass
    def Q(self, parts, marks=None): QUESTIONS.append([])
    def keep_with_next(self): pass
    def solution_box(self, rows, keep_together=True): QUESTIONS[-1].extend(rows)
    def save(self, path): pass
wl.Worksheet = Recorder
BUILD = sys.argv[1]
src = open(BUILD).read().replace('from worksheet_lib import Worksheet', 'from worksheet_lib import Worksheet\nWorksheet = Recorder')
ns = {'Recorder': Recorder, '__name__': 'rec'}
exec(compile(src, BUILD, 'exec'), ns)
split = wl._split_aligned   # patched by build.py (keeps pmatrix rows whole)

FS = 12
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.5)
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(FS)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
pf = st.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.5

t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run(sys.argv[3] if len(sys.argv) > 3 else 'Solutions'); r.bold = True; r.font.size = Pt(14)
t.paragraph_format.space_after = Pt(8)

W = [Cm(1.1), Cm(1.6), Cm(14.3)]
table = doc.add_table(rows=0, cols=3); table.style = 'Table Grid'; table.autofit = False

def set_size(elem, half_pts):
    for rr in elem.iter(f'{{{wl.M_NS}}}r'):
        wrpr = rr.find(qn('w:rPr'))
        if wrpr is None:
            wrpr = OxmlElement('w:rPr'); mrpr = rr.find(f'{{{wl.M_NS}}}rPr')
            (mrpr.addnext(wrpr) if mrpr is not None else rr.insert(0, wrpr))
        if wrpr.find(qn('w:sz')) is None:   # leave the grey 8pt notes alone
            for tag in ('w:sz', 'w:szCs'):
                e = OxmlElement(tag); e.set(qn('w:val'), str(half_pts)); wrpr.append(e)

M = wl.M_NS
def bolden(e):
    for rr in e.iter(f'{{{M}}}r'):
        mrpr = rr.find(f'{{{M}}}rPr')
        if mrpr is None:
            mrpr = OxmlElement('m:rPr'); rr.insert(0, mrpr)
        sty = mrpr.find(f'{{{M}}}sty')
        if sty is None:
            sty = OxmlElement('m:sty'); mrpr.append(sty); sty.set(f'{{{M}}}val', 'bi')
        else:
            v = sty.get(f'{{{M}}}val'); sty.set(f'{{{M}}}val', {'p': 'b', 'i': 'bi'}.get(v, v))
def para_math(p, e):
    mp = OxmlElement('m:oMathPara'); pr = OxmlElement('m:oMathParaPr'); jc = OxmlElement('m:jc')
    jc.set(f'{{{M}}}val', 'left'); pr.append(jc); mp.append(pr); mp.append(e); p._element.append(mp)
EM_CM = FS * 2.54 / 72
from lxml import etree as _et
_orig_em = wl._omml_em
def _em(el):
    tag = _et.QName(el).localname
    if tag == 'm':        # a column vector is as wide as its widest entry, not the sum of its rows
        rows = [r for r in el if _et.QName(r).localname == 'mr']
        return max((sum(_em(c) for c in r) for r in rows), default=0) + 0.3
    if tag == 'acc':      # an arrow over letters: the letters
        e = el.find(f'{{{wl.M_NS}}}e'); return _em(e) if e is not None else 0
    if tag == 't':
        return _orig_em(el)
    Mq = lambda n: el.find(f'{{{wl.M_NS}}}{n}')
    w = lambda x: _em(x) if x is not None else 0
    if tag == 'f': return max(w(Mq('num')), w(Mq('den'))) + 0.15
    if tag in ('sSup', 'sSub'): return w(Mq('e')) + 0.7 * w(Mq('sup' if tag == 'sSup' else 'sub'))
    if tag == 'sSubSup': return w(Mq('e')) + 0.7 * max(w(Mq('sub')), w(Mq('sup')))
    if tag == 'rad': return 0.9 + w(Mq('e'))
    if tag == 'd': return 0.9 + sum(_em(c) for c in el)
    if tag.endswith('Pr') or tag in ('ctrlPr', 'degHide'): return 0.0
    return sum(_em(c) for c in el)
def width_cm(latex):
    e = wl._latex_to_omml(latex, display=False)
    return _em(e) * EM_CM if e is not None else 0
def math_elem(latex, bold=False, colour=None):
    e = wl._latex_to_omml(latex, display=False)
    if e is None: raise ValueError(latex)
    wl._style_annotations(e)
    if colour: wl._colour_math(e, colour)
    if bold: bolden(e)
    set_size(e, FS * 2)
    return e

def new_para(cell, first):
    p = cell.paragraphs[0] if first[0] else cell.add_paragraph()
    first[0] = False
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_together = True
    return p

def fill(p, parts, bold=False, colour=None):
    for part in parts:
        kind, s = part[0], part[1]
        at = part[2] if len(part) > 2 else {}
        b = bold or at.get('bold'); c = colour or (wl._hex(at['color']) if at.get('color') else None)
        if kind == 'text':
            run = p.add_run(s); run.font.size = Pt(FS); run.bold = bool(b)
            if c: run.font.color.rgb = RGBColor.from_string(c)
        else:
            p._element.append(math_elem(s, bold=b, colour=c))

def eq_lines(cell, first, rows, indent_cont=True):
    ind = 0.0
    for lhs, rhs, note in rows:
        if not lhs and not rhs: continue
        p = new_para(cell, first)
        tex = (lhs + ' ' + rhs).strip()
        if note: tex += r' \qquad ' + note
        if lhs: ind = width_cm(lhs)        # continuation lines start under this line's "="
        elif ind: p.paragraph_format.left_indent = Cm(ind)
        para_math(p, math_elem(tex))

def step(cell, first, s):
    if isinstance(s, str):
        rows = split(s)
        if rows is None: rows = [(s, '', '')]
        eq_lines(cell, first, rows)
    elif isinstance(s, tuple) and s[0] == 'or':
        eq_lines(cell, first, split(s[1]))
        p = new_para(cell, first); fill(p, [('text', 'or')])
        eq_lines(cell, first, split(s[2]))
    elif isinstance(s, tuple) and s[0] == 'check':
        p = new_para(cell, first)
        fill(p, [('text', '✓ Check: ', {'bold': True})] + list(s[1]), colour=wl.CHECK_GREEN)
    elif isinstance(s, list):
        p = new_para(cell, first)
        if s and s[0][0] == 'text' and s[0][1] == 'Answer: ':
            fill(p, [('text', '∴ ')] + list(s[1:]), bold=True)
        else:
            fill(p, s)
    else:
        raise ValueError(s)

for qi, rows in enumerate(QUESTIONS, 1):
    for ri, (label, steps) in enumerate(rows):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        for i, w in enumerate(W): row.cells[i].width = w
        if isinstance(label, tuple): lab = label[0] + label[1]
        else: lab = label
        row.cells[0].paragraphs[0].add_run(f'{qi}.' if ri == 0 else '')
        row.cells[1].paragraphs[0].add_run(lab)
        cell = row.cells[2]; first = [True]
        for s in steps: step(cell, first, s)
        # a little air at the foot of each row
        cell.paragraphs[-1].paragraph_format.space_after = Pt(4)
        for c in row.cells[:2]: c.paragraphs[0].paragraph_format.space_before = Pt(2)
        cell.paragraphs[0].paragraph_format.space_before = Pt(2)

# column widths in the grid too
grid = table._tbl.tblGrid
for gc, w in zip(grid.findall(qn('w:gridCol')), W): gc.set(qn('w:w'), str(int(w.twips if hasattr(w,'twips') else w/635)))
out = sys.argv[2]
doc.save(out)
import zipfile
z = zipfile.ZipFile(out).read('word/document.xml').decode()
print('saved', out, 'questions', len(QUESTIONS), 'rows', len(table.rows), 'omath', z.count('<m:oMath>'))
