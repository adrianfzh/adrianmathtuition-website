"""The tile grid in a docx, the way a worksheet would carry it.

The picture goes in the SECTION HEADER as a floating anchor behind the text, so
one paragraph covers every page of a sheet however long it runs, and any PDF
export keeps it.  The picture is the full page, positioned against the PAGE and
not the margin box, so the grid runs edge to edge exactly as RGS's does.
"""
import sys
from docx import Document
from docx.shared import Pt, Cm, Emu, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import qn


def watermark(header, png, width_cm=21.0, height_cm=29.7):
    p = header.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(png, width=Cm(width_cm), height=Cm(height_cm))
    inline = run._r.find(qn('w:drawing'))[0]
    cx, cy = Emu(Cm(width_cm)), Emu(Cm(height_cm))
    graphic = inline.find(qn('a:graphic'))
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="2" '
        'behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:docPr id="91" name="watermark"/>'
        '<wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.append(graphic)
    d = run._r.find(qn('w:drawing'))
    d.remove(inline); d.append(anchor)


png, out, label = sys.argv[1], sys.argv[2], sys.argv[3]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.2)
s.top_margin = s.bottom_margin = Cm(2.0)
doc.styles['Normal'].font.name = 'Georgia'
doc.styles['Normal'].font.size = Pt(11)
watermark(s.header, png)
f = s.footer.paragraphs[0]
f.text = 'AdrianMath Tuition · prepared for Sim Ze Kai · 15 Sep 2026 · AM-7F3K'
f.runs[0].font.size = Pt(7.5)
f.runs[0].font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)
f.alignment = 1
doc.add_paragraph(label).runs[0].bold = True
for ln in ['', 'Solve 2 sin 3x + 1 = 0 for 0° ≤ x ≤ 360°.', '',
           '    sin 3x = −1/2', '    basic angle = 30°',
           '    3x = 210°, 330°, 570°, 690°, 930°, 1050°', '',
           '    x = 70°, 110°, 190°, 230°, 310°, 350°', '',
           'A student writes their working across these lines in pencil, so the mark',
           'has to stay light enough to write over and light enough to photocopy.',
           ''] + [''] * 22:
    doc.add_paragraph(ln)
doc.save(out)
print('wrote', out)
