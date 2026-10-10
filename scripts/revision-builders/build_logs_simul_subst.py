"""A Math · Logarithms practice — the "make simultaneous" type and the "substitution" type.
Adrian, 10 Oct 2026: "could I have a logs ws like those that are like make simultaneous
and substitution type". Every question is a real school prelim question from the bank
(the part named below, word for word); the source never prints on the sheet.
Answers: scripts/revision-builders/verify_logs_simul_subst.py (every original equation
scanned for all real roots; rejected values shown to be undefined).

  1  Raffles Girls 2011 P1 Q3          a33ee202
  2  Juying 2022 P1 Q3                 d1c92354
  3  Greendale 2018 P1 Q5              9c02d4c4   (x = 0 rejected)
  4  VCA (IP) 2021 P2 Q3               2faf427a   (x = -2 rejected)
  5  Orchid Park 2024 P2 Q7            8c663a21   (bank key keeps x = -2/5; log2 x is undefined there — rejected here)
  6  Nan Chiau 2022 P1 Q5              6ad594b6
  7  Maris Stella 2025 P2 Q5(c)        0bdb5b2b   (2^x = -2 rejected)
  8  Hougang 2022 P1 Q13(a)            7a38d494
  9  Jurongville 2025 P2 Q7(b)         90b9db6b
 10  Hua Yi 2023 P2 Q7(a)              085ddbbf   ((sqrt 2)^x = -3 rejected)
 11  Cedar Girls 2025 P1 Q10           78e7b70c   (x = 1 rejected: base of a log)
 12  Raffles Girls 2021 P1 Q2          dbd7007e
 13  Yishun Town 2024 P1 Q6            b833b080
 14  Fairfield Methodist 2022 P2 Q7(a) d1a4ecc4   (y = 0 rejected)   } added 10 Oct 2026, Adrian:
 15  Bukit Merah 2016 P1 Q2            c075dc1f   (x = -2 rejected)  } "3 more tougher simultaneous
 16  Swiss Cottage 2017 P2 Q7          462237c7                      }  questions involving indices and logarithms"
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / '.claude/skills/create-worksheet'))
from worksheet_lib import Worksheet
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH

T = lambda s: ('text', s)
M = lambda s: ('math', s)
D = lambda s: ('math_display', s)
ws = Worksheet()
ws.title('Logarithms — Simultaneous and Substitution Types')
ws.subtitle('Additional Mathematics · Practice')


def eqs(lines, marks=None, indent=1.6, tail=None):
    """The question's equation(s), each on a line of its own at full size, with the
    marks level with the last one. Inline, Word shrinks the index fractions until
    they cannot be read; a display line cannot share its paragraph with the [n]
    tab — so the equations sit in the left cell of a borderless row and the marks
    in the right cell, whose edge is the 15.5 cm marks column."""
    prev = ws.doc.paragraphs[-1]
    prev.paragraph_format.keep_with_next = True
    table = ws.doc.add_table(rows=1, cols=2)
    table.autofit = False
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{side}'); el.set(qn('w:val'), 'nil'); b.append(el)
    table._tbl.tblPr.append(b)
    widths = (14.2, 1.72)   # right edge of the [n] = the 15.5 cm marks tab
    for col, w in zip(table.columns, widths):
        col.width = Cm(w)
    left, right = table.rows[0].cells
    left.width, right.width = Cm(widths[0]), Cm(widths[1])
    for i, latex in enumerate(lines):
        p = left.paragraphs[0] if i == 0 else left.add_paragraph()
        ws._solution_step(p, latex, width=widths[0] - indent)
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.keep_with_next = True
    if tail:
        p = left.add_paragraph()
        ws._fill(p, tail)
        p.paragraph_format.left_indent = Cm(indent - 0.6)
        p.paragraph_format.keep_with_next = True
    right.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.keep_with_next = True
    if marks is not None:
        ws._fill(rp, [T(f'[{marks}]')])
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    return table


def simul(stem, eq1, eq2, marks):
    ws.Q(stem)
    eqs([eq1, eq2], marks=marks)


def solve(eq, marks, stem='Solve the equation'):
    ws.Q([T(stem)])
    eqs([eq], marks=marks)


ws.section('Section A — Two equations: make them simultaneous')
simul([T('Solve the simultaneous equations')],
      r'3^{x} = 27\left(3^{y}\right)', r'\lg(x + 2y) = \lg 5 + \lg 3', 4)
ws.ans([M('x = 7'), T(', '), M('y = 4')])

simul([T('Solve the simultaneous equations')],
      r'3^{y} = \dfrac{1}{3}\left(3^{x - 7}\right)', r'\log_{9}(14 - 2y) - \log_{9}x = 0.5', 5)
ws.ans([M('x = 6'), T(', '), M('y = -2')])

simul([T('Solve, for '), M('x'), T(' and '), M('y'), T(', the simultaneous equations')],
      r'\dfrac{5^{2x}}{5^{3y}} = \dfrac{5}{5^{3}\left(5^{y}\right)}',
      r'\log_{3}(x - 4) = \log_{3}(y - 1) - \log_{3}x', 7)
ws.ans([M('x = 5'), T(', '), M('y = 6')])

simul([T('Solve algebraically the simultaneous equations')],
      r'\lg(8x + 22) - \lg(y - 1) = \dfrac{1}{\log_{2}10}', r'e^{\ln y} = 10^{2\lg x}', 5)
ws.ans([M('x = 6'), T(', '), M('y = 36')])

ws.Q([T('Two simultaneous equations are given:')])
eqs([r'64^{x} \times 4^{y} = 1', r'\log_{\frac{1}{2}}(x - 3y - 6) + 2 = \log_{2}x'])
ws.SQ([T('Show that the second equation can be written as')])
eqs([r'\log_{2}(x - 3y - 6) = \log_{2}\dfrac{4}{x}'], marks=3, indent=2.6)
ws.SQ([T('Solve, for '), M('x'), T(' and '), M('y'), T('.')], marks=7)
ws.ans([T('(a) shown; (b) '), M('x = 1'), T(', '), M('y = -3')])

ws.section('Section B — One equation: use a substitution')
solve(r'\left(\log_{5}y\right)^{2} + \log_{5}\dfrac{1}{y^{3}} = 28', 4)
ws.ans([M('y = 78\\,125'), T(' or '), M(r'y = \dfrac{1}{625}')])

solve(r'4^{x + 1} + 7\left(2^{x}\right) = 2', 4)
ws.ans([M('x = -2')])

solve(r'(\ln x)^{2} + \dfrac{2}{\log_{x}e} = 3', 4)
ws.ans([M('x = e'), T(' or '), M(r'x = e^{-3}')])

solve(r'\log_{3}3x^{2} + 4 = \log_{x}27', 5)
ws.ans([M(r'x = \sqrt{3}'), T(' or '), M(r'x = \dfrac{1}{27}')])

solve(r'2^{x} - \left(\sqrt{2}\right)^{x + 2} = 15', 5, stem='Using a suitable substitution, solve the equation')
ws.ans([M('x = 4.64')])

ws.Q([T('Solve the equation below, leaving your answer in terms of '), M(r'\sqrt{2}'), T('.')])
eqs([r'\log_{4}x - \dfrac{3}{\log_{x}8} = \left(\log_{2}x\right)^{2}'], marks=5)
ws.ans([M(r'x = \dfrac{1}{\sqrt{2}}')])

ws.section('Section C — Both together')
simul([T('Given that '), M('x, y > 0'), T(', solve the simultaneous equations')],
      r'(\lg x)(\lg y) = 2', r'\lg\sqrt{\dfrac{xy}{10}} = 1', 7)
ws.ans([M('x = 10'), T(', '), M('y = 100'), T(' or '), M('x = 100'), T(', '), M('y = 10')])

ws.Q([])
ws.SQ([T('The curve '), M(r'\dfrac{x}{y} + y = 1'), T(' and the line '), M('x - 2y + 6 = 0'),
       T(' intersect at the points '), M('A'), T(' and '), M('B'), T('. Find the '), M('x'),
       T('-coordinate of '), M('A'), T(' and of '), M('B'), T('.')], marks=3)
ws.SQ([T('Hence solve for the value(s) of '), M('x'), T(' in the simultaneous equations')])
eqs([r'\dfrac{\lg x}{\lg y} + \lg y = 1', r'\lg\dfrac{x}{y^{2}} = -6'], marks=3, indent=2.6)
ws.ans([T('(a) '), M('x = -2'), T(' or '), M('x = -12'), T('; (b) '), M('x = 10^{-2}'), T(' or '), M('x = 10^{-12}')])

ws.section('Section D — More simultaneous questions (tougher)')
simul([T('Solve the simultaneous equations')],
      r'e\sqrt{e^{x}} = e^{2y}', r'\log_{4}(x + 2) = 1 + \log_{2}y', 8)
ws.ans([M('x = 2'), T(', '), M('y = 1')])

simul([T('Solve the simultaneous equations')],
      r'0.5^{x}\left(4^{3y}\right) = 16', r'\log_{4}2x + \log_{4}(x + 3y) = 1', 5)
ws.ans([M(r'x = \dfrac{2}{3}'), T(', '), M(r'y = \dfrac{7}{9}')])

ws.Q([])
ws.SQ([T('Given the equation below, express '), M('b'), T(' as a power of '), M('a'), T('.')])
eqs([r'\dfrac{\left(\log_{a}b\right)^{2}}{\log_{b}a} - 27 = 0'], marks=4, indent=2.6)
ws.SQ([T('Hence, solve, for '), M('a'), T(' and '), M('b'), T(', the simultaneous equations')])
eqs([r'ab = 81', r'\dfrac{\left(\log_{a}b\right)^{2}}{\log_{b}a} - 27 = 0'], marks=3, indent=2.6)
ws.ans([T('(a) '), M('b = a^{3}'), T('; (b) '), M('a = 3'), T(', '), M('b = 27')])

print('glued (cm):', [round(h, 1) for h in ws.glued_heights()])
ws.save(sys.argv[1], strict_maths=True)
