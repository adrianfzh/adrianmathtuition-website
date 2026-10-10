"""Two small helpers for hand-picked practice sheets built on worksheet_lib
(10 Oct 2026, the logs / calculus / circles sheets).

P('Find $x$ when …')   -> a parts list: text outside $…$, maths inside.
eqs(ws, [latex, …], marks=n) -> the question's expression(s), each on a line of
its own at full size with the marks level with the last one. Inline, Word
shrinks fractions until they cannot be read; a display line cannot share its
paragraph with the [n] tab — so the lines sit in the left cell of a borderless
row and the marks in the right cell, whose edge is the 15.5 cm marks column.
"""
import re
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH


def P(s):
    out = []
    for i, bit in enumerate(re.split(r'\$(.+?)\$', s)):
        if bit:
            out.append(('math', bit) if i % 2 else ('text', bit))
    return out


def eqs(ws, lines, marks=None, indent=1.6):
    ws.doc.paragraphs[-1].paragraph_format.keep_with_next = True
    table = ws.doc.add_table(rows=1, cols=2)
    table.autofit = False
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{side}'); el.set(qn('w:val'), 'nil'); b.append(el)
    table._tbl.tblPr.append(b)
    widths = (14.2, 1.72)   # right edge of the [n] = the 15.5 cm marks tab
    for col, w in zip(table.columns, widths):
        col.width = Cm(w)
    left, right = table.rows[0].cells
    left.width, right.width = Cm(widths[0]), Cm(widths[1])
    for i, latex in enumerate(lines):
        p = left.paragraphs[0] if i == 0 else left.add_paragraph()
        ws._solution_step(p, latex, width=widths[0] - indent)
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.keep_with_next = True
    right.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.keep_with_next = True
    if marks is not None:
        ws._fill(rp, [('text', f'[{marks}]')])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    return table
