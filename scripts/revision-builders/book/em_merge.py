# -*- coding: utf-8 -*-
"""Stage 2 (E Math book): stitch the prepped copies into one document."""
import os
from docx import Document
from docxcompose.composer import Composer

B = os.path.dirname(os.path.abspath(__file__))
PREP = os.path.join(B, "prepped")
ORDER = [
 "REV 3 Quadratic Equations.docx", "REV 3 Quadratic Equations (Applications).docx",
 "REV 3 Quadratic Graphs.docx", "REV 3 Linear Inequalities.docx", "REV 3 Indices.docx",
 "REV 3 Standard Form.docx", "REV 3 Coordinate Geometry.docx",
 "REV 3 Graphs of Functions.docx", "REV 3 Graphs on Graph Paper.docx",
 "REV 3 Distance and Speed Time Graphs.docx", "REV 3 Trigonometry.docx",
 "REV 3 Arc Length and Sector Area.docx", "REV 3 Congruency and Similarity.docx",
 "REV 3 Geometrical Properties of Circles.docx",
]
master = Document(os.path.join(PREP, ORDER[0]))
c = Composer(master)
for fn in ORDER[1:]:
    c.append(Document(os.path.join(PREP, fn)))
out = os.path.join(B, "merged.docx")
c.save(out)
d = Document(out)
print("merged ->", out)
print("  paragraphs:", len(d.paragraphs), " sections:", len(d.sections),
      " tables:", len(d.tables), " images:",
      sum(1 for r in d.part.rels.values() if "image" in r.reltype))
