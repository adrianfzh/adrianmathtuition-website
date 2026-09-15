# -*- coding: utf-8 -*-
"""Stage 1 (E Math book): normalised WORKING COPIES of the fourteen sheets.

The originals in Dropbox are never opened for writing (ADRIAN-STYLE s6 -- a
sheet handed over in Dropbox is his file).  Everything here is done on a copy
under the scratchpad.

Unlike the A Math book, EVERY sheet is baked: six of the fourteen carry an
explicit Normal of Times New Roman 9.5 pt at 1.5 lines and eight inherit from
their own docDefaults, and a merged book keeps only ONE Normal (the master's,
which is Quadratic Equations' and is bare).  Baking writes back exactly what
each sheet already resolves to, so each sheet on its own is unchanged and the
merge can no longer reach the setting.
"""
import os, shutil, sys
from docx import Document
from docx.shared import Cm, Emu

B = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, B)
from plib import bake, bake_tabs, title_block, strip_headers   # noqa: E402

SRC, OUT = os.path.join(B, "src"), os.path.join(B, "prepped")
os.makedirs(OUT, exist_ok=True)

# (source file, the title that goes in the contents) -- Adrian's own order,
# with Quadratic Equations (Applications) sitting straight after Quadratic
# Equations as a section of its own.
SECTIONS = [
    ("REV 3 Quadratic Equations.docx",                "Quadratic Equations"),
    ("REV 3 Quadratic Equations (Applications).docx", "Quadratic Equations: Applications"),
    ("REV 3 Quadratic Graphs.docx",                   "Quadratic Graphs"),
    ("REV 3 Linear Inequalities.docx",                "Linear Inequalities"),
    ("REV 3 Indices.docx",                            "Indices"),
    ("REV 3 Standard Form.docx",                      "Standard Form"),
    ("REV 3 Coordinate Geometry.docx",                "Coordinate Geometry"),
    ("REV 3 Graphs of Functions.docx",                "Graphs of Functions"),
    ("REV 3 Graphs on Graph Paper.docx",              "Graphs on Graph Paper"),
    ("REV 3 Distance and Speed Time Graphs.docx",     "Distance and Speed-Time Graphs"),
    ("REV 3 Trigonometry.docx",                       "Trigonometry"),
    ("REV 3 Arc Length and Sector Area.docx",         "Arc Length and Sector Area"),
    ("REV 3 Congruency and Similarity.docx",          "Congruency and Similarity"),
    ("REV 3 Geometrical Properties of Circles.docx",  "Geometrical Properties of Circles"),
]

# The two sheets that carry no title in the body -- their topic name lived only
# in a running head, and the headers all come off (they disagree with each other
# and Standard Form's footer still says "N Level EM Revision: Indices").
NEEDS_TITLE = {
    "REV 3 Indices.docx":            ("O Level E Math Revision", "INDICES"),
    "REV 3 Coordinate Geometry.docx": ("O Level E Math Revision", "COORDINATE GEOMETRY"),
}

log = []
for fn, title in SECTIONS:
    dst = os.path.join(OUT, fn)
    shutil.copy2(os.path.join(SRC, fn), dst)
    note = ["baked %d paras / %d runs" % bake(dst)]
    d = Document(dst)
    strip_headers(d)
    if fn in NEEDS_TITLE:
        title_block(d, *NEEDS_TITLE[fn])
        note.append("title block added")
    d.save(dst)
    d = Document(dst)
    sec = d.sections[0]
    if abs(sec.page_width.cm - 21.0) > 0.2:          # US Letter -> A4
        old_text = sec.page_width - sec.left_margin - sec.right_margin
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        side = Emu(int((Cm(21.0) - old_text) / 2))   # keep the line length exactly
        sec.left_margin = sec.right_margin = side
        note.append("US Letter -> A4, side margins %.2f cm (line length kept)" % side.cm)
        d.save(dst)
    n_tab, own_tab = bake_tabs(dst)
    if n_tab:
        note.append("tab ladder baked at %d twips on %d paragraphs" % (own_tab, n_tab))
    d = Document(dst)
    s = d.sections[0]
    log.append("%-34s %5.1fx%5.1f T%.1f L%.1f R%.1f B%.1f  %s"
               % (title[:34], s.page_width.cm, s.page_height.cm, s.top_margin.cm,
                  s.left_margin.cm, s.right_margin.cm, s.bottom_margin.cm,
                  "; ".join(note)))
print("\n".join(log))
