# -*- coding: utf-8 -*-
"""Stage 3 (E Math book): the title styling, the bookmarks, and the cover page.

Each topic title becomes a real Heading 1 so that BOTH PDF export paths build an
outline from it -- Word goes by outline level, LibreOffice goes by the style's
NAME -- and carries a bookmark the cover links to.  Heading 1 is redefined here
so that "Heading 1" costs nothing visually: it is exactly the line Adrian
already types (Times New Roman, 11 pt, bold, centred, 1.5 line spacing).
"""
import os, copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

B = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(B, "merged.docx")
OUT = os.path.join(B, "Sec 3 E Math Overall Revision (Worked Examples).docx")
TNR = "Times New Roman"

# (the title line as it is typed in the sheet, the label on the contents page)
SECTIONS = [
    ("QUADRATIC EQUATIONS",               "Quadratic Equations"),
    ("QUADRATIC EQUATIONS APPLICATIONS",  "Quadratic Equations: Applications"),
    ("QUADRATIC GRAPHS",                  "Quadratic Graphs"),
    ("Linear Inequalities",               "Linear Inequalities"),
    ("INDICES",                           "Indices"),
    ("STANDARD FORM",                     "Standard Form"),
    ("COORDINATE GEOMETRY",               "Coordinate Geometry"),
    ("Graphs of Functions",               "Graphs of Functions"),
    ("Graphs on Graph Paper",             "Graphs on Graph Paper"),
    ("Distance and Speed-Time Graphs",    "Distance and Speed-Time Graphs"),
    ("TRIGONOMETRY",                      "Trigonometry"),
    ("ARC LENGTH AND SECTOR AREA",        "Arc Length and Sector Area"),
    ("CONGRUENCY AND SIMILARITY",         "Congruency and Similarity"),
    ("Geometrical Properties of Circles", "Geometrical Properties of Circles"),
]


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), v)
    return e


def rpr(size=22, bold=True, italic=False, color=None):
    r = OxmlElement('w:rPr')
    r.append(el('w:rFonts', ascii=TNR, hAnsi=TNR, eastAsia=TNR, cs=TNR))
    if bold:   r.append(OxmlElement('w:b'))
    if italic: r.append(OxmlElement('w:i'))
    if color:  r.append(el('w:color', val=color))
    r.append(el('w:sz', val=str(size)))
    r.append(el('w:szCs', val=str(size)))
    return r


def run(text, **kw):
    r = OxmlElement('w:r')
    r.append(rpr(**kw))
    t = OxmlElement('w:t'); t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    return r


def para(jc='center', line='360', after=None, before=None, ind=None):
    p = OxmlElement('w:p')
    ppr = OxmlElement('w:pPr')
    sp = el('w:spacing', line=line, lineRule='auto')
    if after is not None:  sp.set(qn('w:after'), after)
    if before is not None: sp.set(qn('w:before'), before)
    ppr.append(sp)
    if ind is not None:
        ppr.append(el('w:ind', left=ind))
    ppr.append(el('w:jc', val=jc))
    p.append(ppr)
    return p


doc = Document(SRC)
body = doc.element.body

# ---------------------------------------------------------------- Heading 1
styles = doc.styles.element
for st in styles.findall(qn('w:style')):
    if st.get(qn('w:styleId')) == 'Heading1':
        styles.remove(st)
h = el('w:style', type='paragraph', styleId='Heading1')
h.append(el('w:name', val='heading 1'))
h.append(el('w:basedOn', val='Normal'))
h.append(el('w:next', val='Normal'))
h.append(OxmlElement('w:qFormat'))
hp = OxmlElement('w:pPr')
hp.append(OxmlElement('w:keepNext'))
hp.append(el('w:spacing', line='360', lineRule='auto'))
hp.append(el('w:outlineLvl', val='0'))
hp.append(el('w:jc', val='center'))
h.append(hp)
h.append(rpr(22))
styles.append(h)

# ------------------------------------------------- the 14 topic-title lines
# Matched IN ORDER, each from where the previous one was found: five of the
# sheets repeat a word that also names another section, and a plain search would
# pick the wrong paragraph.
paras = doc.paragraphs
titles, at = [], 0
for want, _ in SECTIONS:
    for i in range(at, len(paras)):
        if paras[i].text.strip() == want:
            titles.append(paras[i]); at = i + 1; break
    else:
        raise AssertionError("no title line %r after paragraph %d" % (want, at))
print("  titles at paragraphs:", [paras.index(p) for p in titles])

next_id = 1 + max([int(b.get(qn('w:id')) or 0)
                   for b in body.iter(qn('w:bookmarkStart'))] or [0])
contents = []
for n, (p, (_, label)) in enumerate(zip(titles, SECTIONS), 1):
    name = "sec%02d" % n
    contents.append((name, label))
    ppr = p._p.find(qn('w:pPr'))
    if ppr is None:
        ppr = OxmlElement('w:pPr'); p._p.insert(0, ppr)
    for old in ppr.findall(qn('w:pStyle')):
        ppr.remove(old)
    ppr.insert(0, el('w:pStyle', val='Heading1'))
    # one look for all fourteen: several sheets left the title's size to inherit
    for r in p._p.findall(qn('w:r')):
        old = r.find(qn('w:rPr'))
        if old is not None:
            r.remove(old)
        r.insert(0, rpr(22))
    bs = el('w:bookmarkStart', id=str(next_id), name=name)
    be = el('w:bookmarkEnd', id=str(next_id))
    next_id += 1
    p._p.insert(1, bs)          # after w:pPr
    p._p.append(be)

# ------------------------------------------------------------- the cover page
first = body[0]
def before(e):
    first.addprevious(e)

t = para(after='0'); t.append(run('Sec 3 E Math', size=36)); before(t)
t = para(after='0'); t.append(run('Overall Revision', size=36)); before(t)
t = para(after='240'); t.append(run('Worked Examples', size=26)); before(t)
t = para(jc='left', after='120'); t.append(run('Contents', size=24)); before(t)
for name, label in contents:
    p = para(jc='left', after='0', ind='284')
    hl = OxmlElement('w:hyperlink')
    hl.set(qn('w:anchor'), name)
    # black, not underlined -- a printed page should not look like a web page
    hl.append(run(label, size=22, bold=False, color='000000'))
    p.append(hl)
    before(p)
t = para(jc='left', before='240', after='0')
t.append(run('Tap a topic to jump to it.', size=20, bold=False,
             italic=True, color='595959'))
before(t)

# the cover stands on its own page
ppr = first.find(qn('w:pPr'))
if ppr is None:
    ppr = OxmlElement('w:pPr'); first.insert(0, ppr)
if ppr.find(qn('w:pageBreakBefore')) is None:
    ppr.insert(0, OxmlElement('w:pageBreakBefore'))

# ---------------------------------------------- the document's own last section
# docxcompose keeps the MASTER's body sectPr, which is Quadratic Equations'.
# The final section of the book is Circles, so its page setup belongs here.
last_sheet = Document(os.path.join(B, "prepped",
                                   "REV 3 Geometrical Properties of Circles.docx"))
old = body.find(qn('w:sectPr'))
body.replace(old, copy.deepcopy(last_sheet.element.body.find(qn('w:sectPr'))))

# ---------------------------------------------------------------- tab stops
# One settings.xml governs the whole book.  prep.py baked each sheet's own tab
# ladder onto its paragraphs wherever that ladder was not 720 twips, so the
# book's default is set to 720 to agree with them.
st = doc.settings.element
d = st.find(qn('w:defaultTabStop'))
if d is None:
    d = OxmlElement('w:defaultTabStop'); st.append(d)
d.set(qn('w:val'), '720')

# ------------------------------------------------------ Word's layout engine
# Nine of the fourteen sheets were written in compatibility mode 14 and five in
# 15, and one settings.xml governs the whole book.  Mode 15 re-wraps the mode-14
# sheets and grows them (A Math Circles 6 pages -> 8, measured), so the book is
# set to 14 -- the mode the majority were laid out in, and the one that leaves
# Adrian's page breaks where he put them.
for c in st.iter(qn('w:compatSetting')):
    if c.get(qn('w:name')) == 'compatibilityMode':
        c.set(qn('w:val'), '14')

# ------------------------------------------------ the paragraph Word adds back
# A Word document may not end on a table: Word supplies an empty paragraph after
# it at Normal's size, and if the last table finishes close to the foot of the
# page that paragraph opens a BLANK final page.  Writing it explicitly, one
# point tall with no spacing, keeps the document legal and keeps it on the page
# the table ends on.
last = OxmlElement('w:p')
lppr = OxmlElement('w:pPr')
sp = OxmlElement('w:spacing')
sp.set(qn('w:after'), '0'); sp.set(qn('w:line'), '20'); sp.set(qn('w:lineRule'), 'exact')
lppr.append(sp)
lrpr = OxmlElement('w:rPr')
for tag, val in (('w:sz', '2'), ('w:szCs', '2')):
    e = OxmlElement(tag); e.set(qn('w:val'), val); lrpr.append(e)
lppr.append(lrpr)
last.append(lppr)
body.find(qn('w:sectPr')).addprevious(last)

doc.save(OUT)

d = Document(OUT)
print("wrote", OUT)
print("  paragraphs:", len(d.paragraphs), " tables:", len(d.tables),
      " sections:", len(d.sections))
print("  bookmarks:", [b.get(qn('w:name')) for b in d.element.body.iter(qn('w:bookmarkStart'))])
print("  anchors  :", [a.get(qn('w:anchor')) for a in d.element.body.iter(qn('w:hyperlink'))])
print("  headings :", [p.text for p in d.paragraphs if p.style.name == 'Heading 1'])
