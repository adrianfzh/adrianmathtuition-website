# -*- coding: utf-8 -*-
"""Stage 1b (E Math book): cut the Practice blocks out.

Adrian, 15 Sep 2026: "just the worked examples will do" -- the same cut the A
Math book took ("cut trig practice for consistency").

The cuts are given as BODY-CHILD INDICES read off the source sheets rather than
as heading text.  Five of the fourteen sheets carry the word "Practice" two to
four times, and one of the stop texts ("Notes:") first occurs BEFORE its own
Practice block -- a text search would have found that first occurrence and cut
nothing.  Every index is re-checked against the text that must be sitting there
before anything is removed, so a changed source fails loudly instead of quietly
cutting the wrong thing.

Indices and Coordinate Geometry are +2 on their source indices: prep.py put a
two-line title block in front of them.
"""
import os, sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

B = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, B)
from clib import break_onto_next, drop_before_break, drop_trailing_breaks  # noqa: E402

P = os.path.join(B, "prepped")
E = None                       # "to the end of the sheet"

# file -> [(start, stop, the text that must be at start, the text at stop)]
CUTS = {
 "REV 3 Quadratic Equations.docx":                [(41, E, "Practice", None)],
 "REV 3 Quadratic Equations (Applications).docx": [(51, E, "Practice", None)],
 "REV 3 Quadratic Graphs.docx":                   [(41, E, "Practice", None)],
 "REV 3 Linear Inequalities.docx": [
     (62,  97, "Practice", "Section B"),
     (138, 154, "Practice", "Section C"),
     (188, E,  "Practice", None)],
 "REV 3 Indices.docx":       [(81, E, "Indices Practice", None)],      # +2
 "REV 3 Standard Form.docx": [(34, E, "Practice", None)],
 "REV 3 Coordinate Geometry.docx": [                                    # +2
     (51,  95, "Practice", "Notes:"),
     (115, E,  "Practice", None)],
 "REV 3 Graphs of Functions.docx": [
     (48,  97,  "Practice", "Section B"),
     (157, 188, "Practice", "Section C"),
     (227, E,   "Practice", None)],
 "REV 3 Graphs on Graph Paper.docx": [
     (46,  78,  "Practice", "Section B"),
     (121, 169, "Practice", "Section C"),
     (224, 278, "Practice", "Section D"),
     (334, E,   "Practice", None)],
 "REV 3 Distance and Speed Time Graphs.docx": [
     (49,  67,  "Practice", "Section B"),
     (125, 183, "Practice", "Section C"),
     (225, E,   "Practice", None)],
 "REV 3 Trigonometry.docx":               [(51, E, "Practice", None)],
 "REV 3 Arc Length and Sector Area.docx": [(59, E, "Practice", None)],
 "REV 3 Congruency and Similarity.docx":  [(70, E, "Congruency and Similarity Practice", None)],
 "REV 3 Geometrical Properties of Circles.docx": [(61, E, "Practice", None)],
}

ORDER = [
 "REV 3 Quadratic Equations.docx", "REV 3 Quadratic Equations (Applications).docx",
 "REV 3 Quadratic Graphs.docx", "REV 3 Linear Inequalities.docx", "REV 3 Indices.docx",
 "REV 3 Standard Form.docx", "REV 3 Coordinate Geometry.docx",
 "REV 3 Graphs of Functions.docx", "REV 3 Graphs on Graph Paper.docx",
 "REV 3 Distance and Speed Time Graphs.docx", "REV 3 Trigonometry.docx",
 "REV 3 Arc Length and Sector Area.docx", "REV 3 Congruency and Similarity.docx",
 "REV 3 Geometrical Properties of Circles.docx",
]


def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn('w:t'))).strip()


for k, fn in enumerate(ORDER):
    path = os.path.join(P, fn)
    doc = Document(path)
    body = doc.element.body
    msg = []
    kids = [e for e in body if e.tag != qn('w:sectPr')]
    # check every boundary BEFORE removing anything
    for start, stop, want_a, want_b in CUTS[fn]:
        got = text_of(kids[start])
        assert got == want_a, "%s: index %d is %r, expected %r" % (fn, start, got[:60], want_a)
        if stop is not E:
            got = text_of(kids[stop])
            assert got.startswith(want_b), \
                "%s: index %d is %r, expected %r..." % (fn, stop, got[:60], want_b)
    # remove from the BOTTOM up so the earlier indices stay true
    for start, stop, _, _ in sorted(CUTS[fn], reverse=True):
        j = len(kids) if stop is E else stop
        for el in kids[start:j]:
            body.remove(el)
        msg.append("cut %d blocks at [%d:%s]" % (j - start, start, "end" if stop is E else stop))
    drop_trailing_breaks(body, msg)
    # a page break of its own is carried by the block after it
    break_onto_next(doc, body, msg)
    # and an empty paragraph in front of one is dead weight
    drop_before_break(doc, body, msg)
    # every sheet after the first starts on a new page
    if k > 0:
        first = doc.paragraphs[0]._p
        ppr = first.find(qn('w:pPr'))
        if ppr is None:
            ppr = OxmlElement('w:pPr'); first.insert(0, ppr)
        if ppr.find(qn('w:pageBreakBefore')) is None:
            ppr.insert(0, OxmlElement('w:pageBreakBefore'))
        msg.append("starts on a new page")
    doc.save(path)
    d = Document(path)
    print("%-46s %4d paras %3d tables  %s"
          % (fn[:46], len(d.paragraphs), len(d.tables), "; ".join(msg)))
