#!/usr/bin/env python3
"""GCE paper JSON → Word, in Adrian's house style.

  python3 scripts/gce-paper/export-docx.py data/gce-generated/GCE-AM-P1-seed1-2026-09-08.json \
      --figures <dir holding Q<n>.png> --out <dir> [--space 3.0]

Writes <name>.docx (the paper: front page, formulae, questions with writing
space) and <name>-solutions.docx (every question with its boxed worked solution,
then an answer key). Everything comes from the assembled JSON — no model calls.
Typesetting goes through .claude/skills/create-worksheet/worksheet_lib.py
(python-docx + pandoc for native Word equations), so the output matches the
worksheets Adrian already edits.
"""
import argparse
import json
import os
import re
import sys
from os.path import abspath, basename, dirname, exists, join, splitext

ROOT = dirname(dirname(dirname(abspath(__file__))))
sys.path.insert(0, join(ROOT, '.claude', 'skills', 'create-worksheet'))

import worksheet_lib  # noqa: E402
from worksheet_lib import Worksheet  # noqa: E402
from docx.shared import Cm  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

# ---------------------------------------------------------------- maths ----

_orig_omml = worksheet_lib._latex_to_omml
_omml_cache = {}
FAILED = []


def _memo_omml(latex, display=False):
    key = (latex, bool(display))
    if key not in _omml_cache:
        try:
            _omml_cache[key] = _orig_omml(latex, display)
        except Exception as e:  # pandoc refused it — keep the LaTeX visible, never drop it
            FAILED.append((latex, str(e)[:80]))
            _omml_cache[key] = None
    el = _omml_cache[key]
    if el is None:
        return None
    # pandoc hands back one lxml element; a second use in the document needs a copy
    import copy
    return copy.deepcopy(el)


worksheet_lib._latex_to_omml = _memo_omml

# An escaped \$ (a price) never opens or closes a maths run.
MATH_RE = re.compile(r'\$\$(.+?)\$\$|\\\[(.+?)\\\]|(?<!\\)\$((?:\\.|[^$\\])+?)\$', re.S)


DEG_RE = re.compile(r'\^\s*\{\s*\\circ\s*\}|\^\s*\\circ')
TABLE_ROW_RE = re.compile(r'^\s*\|.*\|\s*$')
TABLE_RULE_RE = re.compile(r'^\s*\|[\s:|-]+\|\s*$')
# A line that is one LaTeX array: `$\begin{array}{|c|c|} \hline … \end{array}$`.
ARRAY_LINE_RE = re.compile(
    r'^\s*\$\$?\s*\\begin\{array\}\{[^}]*\}(.*?)\\end\{array\}\s*\$?\$\s*\.?\s*$', re.S)
TEXT_CMD_RE = re.compile(r'\\text\{([^{}]*)\}')
PLAIN_CELL_RE = re.compile(r'^(?:\\\$|[\d\s.,:%$-])+$')
LATEX_SPACE_RE = re.compile(r'\\[;,:!]|\\q?quad\b')


def array_cell(cell):
    """One LaTeX array cell as a segs() string: \\text{…} runs become prose, a
    bare number or price stays prose, anything else becomes one $…$ run."""
    out, pos = [], 0

    def maths(s):
        if not s.strip():
            out.append(s)
        elif PLAIN_CELL_RE.match(LATEX_SPACE_RE.sub(' ', s).strip()):
            out.append(LATEX_SPACE_RE.sub(' ', s))
        else:
            out.append('$' + s.strip() + '$')

    for m in TEXT_CMD_RE.finditer(cell):
        maths(cell[pos:m.start()])
        out.append(m.group(1))
        pos = m.end()
    maths(cell[pos:])
    return re.sub(r'\s+', ' ', ''.join(out)).strip()


def array_rows(body):
    """A LaTeX array body as table rows (the \\hline rules dropped: a data
    table is drawn with all its rules). An empty cell stays empty."""
    rows = []
    for line in body.replace('\\hline', ' ').split('\\\\'):
        cells = [array_cell(c.strip()) for c in line.split('&')]
        if any(cells):
            rows.append(cells)
    return rows


def _clean(s):
    return s.replace('**', '').replace('\\$', '$')


def segs(text, attrs=None):
    """Split '$…$' / '$$…$$' / '\\[…\\]' out of a line into worksheet_lib parts."""
    out, pos = [], 0
    text = text or ''

    def txt(s):
        if not s:
            return
        out.append(('text', _clean(s), attrs) if attrs else ('text', _clean(s)))

    for m in MATH_RE.finditer(text):
        before = text[pos:m.start()]
        if m.group(3) is not None:
            latex = DEG_RE.sub('°', m.group(3).strip())
            # "cm$^{2}$": a unit's power typed as a bare superscript has no base
            # (Word draws an empty box) — the unit goes into the maths as its base.
            unit = re.search(r'([A-Za-z]+)$', before)
            if latex.startswith('^') and unit:
                before = before[:unit.start()]
                latex = '\\text{' + unit.group(1) + '}' + latex
        txt(before)
        if m.group(3) is not None:
            if _memo_omml(latex, False) is None:
                out.append(('text', latex, {'italic': True}))
            else:
                out.append(('math', latex))
        else:
            latex = (m.group(1) or m.group(2)).strip()
            if _memo_omml(latex, True) is None:
                out.append(('text', latex, {'italic': True}))
            else:
                out.append(('math_display', latex))
        pos = m.end()
    txt(text[pos:])
    return out


def whole_math(line):
    """A line that is exactly one $…$ (optionally ending in a full stop) is a display step."""
    m = re.fullmatch(r'\s*\$(.+?)\$\s*\.?\s*', line, re.S)
    if m and '$' not in m.group(1):
        return m.group(1).strip()
    return None


# --------------------------------------------------------------- layout ----

Q_TEXT_CM = 1.0   # numbering.xml puts question text at 567 twips
STEP_CM = 0.9     # each part level steps in by this much

LABEL_RE = re.compile(r'^\s*\(([a-z]{1,3})\)\s*(?:\(([ivx]{1,5})\))?\s*', re.I)
ROMAN = {'i', 'ii', 'iii', 'iv', 'v', 'vi'}


def blocks(text):
    """A text's lines, with a markdown pipe table or a LaTeX array line folded
    into one ('table', rows) block and a display line of items spaced by
    \\qquad into ('spaced', items). A LaTeX array left as maths reached Word as
    a matrix with no rules, its empty cells and \\hline as red ¿ in the
    LibreOffice PDF (23 Sep 2026, E Math Set 2 P2 Q4, Q7, Q9 — the fix the
    revision builders got on 18 Sep, build_lib.stem_blocks)."""
    out, rows = [], []
    for line in (text or '').split('\n'):
        if TABLE_ROW_RE.match(line):
            if not TABLE_RULE_RE.match(line):
                rows.append([c.strip() for c in line.strip().strip('|').split('|')])
            continue
        if rows:
            out.append(('table', rows)); rows = []
        if not line.strip():
            continue
        arr = ARRAY_LINE_RE.match(line)
        if arr and array_rows(arr.group(1)):
            out.append(('table', array_rows(arr.group(1))))
            continue
        m = re.fullmatch(r'\s*\$\$?([^$]+?)\$?\$\s*', line, re.S)   # a line that is one maths run
        if m and '\\qquad' in m.group(1):
            out.append(('spaced', [x.strip() for x in re.split(r'(?:\\qquad\s*)+', m.group(1)) if x.strip()]))
        else:
            out.append(('line', line))
    if rows:
        out.append(('table', rows))
    return out


def special_block(ws, kind, body):
    if kind == 'table':
        ws.data_table([[segs(c) or [('text', '')] for c in row] for row in body])
    else:
        parts = []
        for i, item in enumerate(body):
            if i:
                parts.append(('text', '\u2003\u2003\u2003'))
            parts.append(('math', DEG_RE.sub('°', item)))
        ws._add(parts, alignment=WD_ALIGN_PARAGRAPH.CENTER)


def split_label(label):
    m = LABEL_RE.match(label or '')
    if not m:
        return (label or '').strip(), ''
    return '(' + m.group(1) + ')', ('(' + m.group(2) + ')') if m.group(2) else ''


def indent(p, level, hang_levels):
    pf = p.paragraph_format
    text_cm = Q_TEXT_CM + STEP_CM * (level + 1)
    pf.left_indent = Cm(text_cm)
    pf.first_line_indent = Cm(-STEP_CM * hang_levels) if hang_levels else Cm(0)
    for lv in range(level + 1):
        pf.tab_stops.add_tab_stop(Cm(Q_TEXT_CM + STEP_CM * (lv + 1)))


def number_line(ws, labels, line, marks, level):
    """A question with no stem: its first part sits on the number's own line."""
    head = [('text', '\t'.join(labels) + '\t')] if any(labels) else []
    p = ws.Q(head + segs(line), marks=marks)
    if any(labels):
        pf = p.paragraph_format
        text_cm = Q_TEXT_CM + STEP_CM * (level + 1)
        pf.left_indent = Cm(text_cm)
        pf.first_line_indent = Cm(-text_cm)
        for lv in range(level + 1):
            pf.tab_stops.add_tab_stop(Cm(Q_TEXT_CM + STEP_CM * lv))
    return p


def labelled(ws, labels, text, marks, level, numbered=False):
    """One part: `labels` is the list of label strings to lay across the hanging
    tab stops (['(c)', '(i)'] or ['', '(ii)'] or ['(a)']), text may span lines."""
    bl = blocks(text) or [('line', '')]
    last_line = max((i for i, b in enumerate(bl) if b[0] == 'line'), default=-1)
    saved_space = ws.working_space
    if bl[-1][0] == 'table':         # a table to complete IS the answer space
        ws.working_space = 0
    try:
        _labelled_blocks(ws, bl, last_line, labels, marks, level, numbered)
    finally:
        ws.working_space = saved_space


def _labelled_blocks(ws, bl, last_line, labels, marks, level, numbered):
    for i, (kind, body) in enumerate(bl):
        if kind != 'line':
            special_block(ws, kind, body)
            continue
        m = marks if i == last_line else None
        if i == 0 and numbered:
            number_line(ws, labels, body, m, level)
        else:
            head = [('text', '\t'.join(labels) + '\t')] if i == 0 and any(labels) else []
            p = ws._add(head + segs(body), marks=m)
            indent(p, level if any(labels) else -1, len(labels) if i == 0 and any(labels) else 0)
        if i + 1 < len(bl):
            ws.keep_with_next()      # a part is never cut across two pages, its table included


def stem_paras(ws, q, marks_on_stem):
    bl = blocks(q.get('stem') or '')
    last_line = max((i for i, b in enumerate(bl) if b[0] == 'line'), default=-1)
    first = True
    for i, (kind, body) in enumerate(bl):
        if kind != 'line':
            special_block(ws, kind, body)
            continue
        m = marks_on_stem if i == last_line else None
        if first:
            ws.Q(segs(body), marks=m)
            first = False
        else:
            p = ws.para(segs(body), marks=m)
            p.paragraph_format.left_indent = Cm(Q_TEXT_CM)
        if i + 1 < len(bl):
            ws.keep_with_next()      # the stem's lines and its tables stay together
    return not first


def asked_width_cm(figures, pos):
    """<run>/figure-sizes.json = {"26": 105} — printed widths in mm Adrian asked
    for; generate.mjs reads the same file for the PDF."""
    try:
        p = join(figures, 'figure-sizes.json') if figures else None
        mm = float(json.load(open(p)).get(str(pos), 0)) if p and exists(p) else 0
        return mm / 10 if mm else None
    except Exception:
        return None


def figure_width_cm(figures, pos, png):
    asked = asked_width_cm(figures, pos)
    if asked:
        return asked
    """The same printed sizes as the PDF (generate.mjs figureDataUri): a drawing
    10 cm wide, 12 cm when wide (aspect >= 1.5), never taller than 10 cm; a
    graph-paper grid at one major square = 1 cm exactly."""
    spec = join(figures, f'Q{pos}.figure.json') if figures else None
    svg = join(figures, f'Q{pos}.figure.svg') if figures else None
    try:
        if spec and exists(spec) and json.load(open(spec)).get('family') == 'graph-paper' and svg and exists(svg):
            text = open(svg).read()
            m = re.search(r'<svg[^>]*\swidth="([\d.]+)"', text)
            grid = re.search(r'<path d="([^"]+)"', text)
            xs = sorted({float(a) for a, b in re.findall(r'M ([\d.]+) [\d.]+ L ([\d.]+) ', grid.group(1)) if a == b})
            minor = min(b - a for a, b in zip(xs, xs[1:]))
            return round(float(m.group(1)) / (minor * 5), 2)
    except Exception:
        pass
    try:
        from PIL import Image
        w, h = Image.open(png).size
        width = 12.0 if w / h >= 1.5 else 10.0
        if h * width / w > 10.0:
            width = 10.0 * w / h
        return round(width, 2)
    except Exception:
        return 10.0


def has_figure(figures, pos):
    return bool(figures) and any(exists(join(figures, n)) for n in (f'Q{pos}.figure.png', f'Q{pos}.png'))


def figure_para(ws, q, pos, figures, raw=False, key=None):
    # figure.mjs writes Q<n>.figure.png beside the draft; a hand-made Q<n>.png also counts.
    # key = "5a": a PER-PART figure (Q5a.figure.png) printed under the part whose
    # "figure": "a" names it (20 Sep 2026).
    path = None
    key = key or str(pos)
    for name in (f'Q{key}.figure.png', f'Q{key}.png'):
        cand = join(figures, name) if figures else None
        if cand and exists(cand):
            path = cand
            break
    if path and raw:
        # an answer space keeps its blank paper: ws.figure trims a PNG to its ink
        para = ws.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(path, width=Cm(figure_width_cm(figures, key, path)))
        ws._block_paras.append(para)
        return True
    if path:
        ws.figure(path, width_cm=figure_width_cm(figures, key, path))
        return True
    if q.get('needs_figure') and key == str(pos):
        p = ws.para([('text', '[Figure to be drawn: ' + (q.get('figure_description') or '') + ']', {'italic': True})])
        p.paragraph_format.left_indent = Cm(Q_TEXT_CM)
    return False


def grid_part(q, figures, pos):
    """A graph-paper grid prints where the paper says "On the grid" — the
    first part (or sub-part) whose text mentions the grid — not above the
    question like a diagram. None for every other figure."""
    spec = join(figures, f'Q{pos}.figure.json') if figures else None
    try:
        if not (spec and exists(spec) and json.load(open(spec)).get('family') == 'graph-paper'):
            return None
    except Exception:
        return None
    for part in q.get('parts') or []:
        if re.search(r'\bgrid\b', part.get('text') or '', re.I):
            return part
        for sub in part.get('subparts') or []:
            if re.search(r'\bgrid\b', sub.get('text') or '', re.I):
                return part
    return None


PAGE_USABLE_CM = 29.7 - 2.0 - 1.0      # worksheet_lib page: A4, 2 cm top, 1 cm bottom
LINE_CM = 0.60                          # one 1.5-spaced text line / one blank writing line
TITLE_CM = 2.6                          # title + subtitle block on the first page


def _plain_len(text):
    return len(re.sub(r'\$[^$]*\$', lambda m: 'x' * max(3, len(m.group(0)) // 2), text or ''))


def _figure_h_cm(figures, key, png):
    try:
        from PIL import Image
        w, h = Image.open(png).size
        return figure_width_cm(figures, key, png) * h / w
    except Exception:
        return 7.0


def estimate_used_cm(q, figures, pos, first_page):
    """Rough printed height of a question WITHOUT its writing space, so the space
    can be sized to fill the rest of the page (layout.page_per_question)."""
    used = TITLE_CM if first_page else 0.4
    def block(text, chars):
        n = _plain_len(text)
        return (max(1, -(-n // chars)) * LINE_CM + 0.15) if n else 0
    used += block(q.get('stem') or '', 88)
    parts = q.get('parts') or []
    for part in parts:
        used += block(part.get('text') or '', 78)
        for sub in part.get('subparts') or []:
            used += block(sub.get('text') or '', 70)
        if part.get('figure'):
            png = join(figures, f"Q{pos}{part['figure']}.figure.png") if figures else None
            if png and exists(png):
                used += _figure_h_cm(figures, f"{pos}{part['figure']}", png) + 0.5
    png = join(figures, f'Q{pos}.figure.png') if figures else None
    if png and exists(png) and not any(p.get('figure') for p in parts):
        used += _figure_h_cm(figures, str(pos), png) + 0.5
    return used


def question(ws, s, figures, with_marks=True):
    q = s.get('question') or s.get('draft')
    parts = q.get('parts') or []
    has_stem = stem_paras(ws, q, s['target'] if (with_marks and not parts) else None)
    # 'answer_space': the figure IS the space the candidate draws in (a
    # construction's given line) — it prints under the last part, and the
    # parts above it leave no writing lines of their own.
    in_answer_space = q.get('figure_position') == 'answer_space' and bool(parts)
    after_part = parts[-1] if in_answer_space else grid_part(q, figures, s['pos'])
    part_figures = any(p.get('figure') for p in parts)   # each part prints its own Q<n><letter> figure
    if after_part is None and not part_figures:
        if has_stem and has_figure(figures, s['pos']):
            for para in ws._block_paras:
                para.paragraph_format.keep_with_next = True
        figure_para(ws, q, s['pos'], figures)
    prev_outer = None
    numbered = not has_stem          # no stem: the first part carries the number
    for part in parts:
        saved_space = ws.working_space
        part_fig = part.get('figure')                       # "a" → Q<n>a.figure.png under this part
        part_raw = part.get('figure_position') == 'answer_space'
        if part is after_part or (in_answer_space and with_marks) or part_fig:
            ws.working_space = 0          # a part with its own figure gets its space AFTER the figure
        outer, inner = split_label(part.get('label', ''))
        subs = part.get('subparts') or []
        marks = part.get('marks') if with_marks else None
        text = part.get('text', '')
        merge_first_sub = bool(subs) and not text.strip()   # "(a)  (i)  …" on one line
        if merge_first_sub:
            pass
        elif inner:
            labels = [outer if outer != prev_outer else '', inner]
            labelled(ws, labels, text, None if subs else marks, level=1, numbered=numbered)
            numbered = False
        else:
            labelled(ws, [outer], text, None if subs else marks, level=0, numbered=numbered)
            numbered = False
        if part_fig and not subs:
            # the part's own figure sits right under its text (before the writing space);
            # an answer-space figure IS the space, otherwise the space follows the figure
            for para in ws._block_paras[-2:]:
                para.paragraph_format.keep_with_next = True
            figure_para(ws, q, s['pos'], figures, raw=part_raw, key=f"{s['pos']}{part_fig}")
            if not part_raw and with_marks and marks and saved_space:
                ws.working_space = saved_space
                ws.workspace(marks=marks)
        prev_outer = outer
        for k, sub in enumerate(subs):
            so, si = split_label(sub.get('label', ''))
            labels = [outer if k == 0 else '', si or so] if merge_first_sub else [si or so]
            labelled(ws, labels, sub.get('text', ''), sub.get('marks') if with_marks else None, level=1,
                     numbered=numbered)
            numbered = False
        ws.working_space = saved_space
        if part is after_part:
            figure_para(ws, q, s['pos'], figures, raw=in_answer_space)
        elif part_fig and subs:
            figure_para(ws, q, s['pos'], figures, raw=part_raw, key=f"{s['pos']}{part_fig}")
    return q


# ----------------------------------------------------------- solutions ----


def split_long_math(latex, limit=70):
    """A display line longer than `limit` chars with 2+ top-level '=' signs
    becomes an aligned block, one '=' per line, so it stacks instead of
    running off the solution box (OMML display math never wraps)."""
    if len(latex) <= limit:
        return latex
    pieces, depth, cur = [], 0, ''
    for ch in latex:
        if ch in '{([':
            depth += 1
        elif ch in '})]':
            depth -= 1
        if ch == '=' and depth == 0:
            pieces.append(cur)
            cur = ''
        else:
            cur += ch
    pieces.append(cur)
    if len(pieces) < 3:
        return latex
    first, rest = pieces[0].strip(), [q.strip() for q in pieces[1:]]
    return '\\begin{aligned} ' + first + ' &= ' + ' \\\\ &= '.join(rest) + ' \\end{aligned}'

def solution_rows(sol):
    rows, label, steps, prev_outer = [], None, [], None

    def flush():
        if label is not None or steps:
            rows.append((label or '', steps or [[('text', '')]]))

    for raw in (sol or '').split('\n'):
        line = raw.strip()
        if not line:
            continue
        m = LABEL_RE.match(line)
        if m and (m.end() < len(line) or not steps):
            g1, g2 = m.group(1).lower(), (m.group(2) or '').lower()
            if g1 in ROMAN and prev_outer:
                new = prev_outer + '(' + g1 + ')'
            else:
                new = '(' + g1 + ')' + (('(' + g2 + ')') if g2 else '')
                prev_outer = '(' + g1 + ')'
            if new != label:
                flush()
                label, steps = new, []
            line = line[m.end():].strip()
            if not line:
                continue
        chk = re.match(r'^\(?\s*Check\s*:\s*(.*?)\)?\s*$', line, re.I | re.S)
        if chk:
            steps.append(('check', segs(chk.group(1))))
            continue
        disp = whole_math(line)
        if disp is not None:
            disp = split_long_math(disp)
        if disp is not None and _memo_omml(disp, True) is not None:
            steps.append(disp)
        else:
            steps.append(segs(line))
    flush()
    return rows or [('', [[('text', '(no worked solution recorded)')]])]


# ------------------------------------------------------------ front page ----

INSTRUCTIONS = [
    'Answer all the questions.',
    'Write your answers in the spaces provided.',
    'If working is needed for any question it must be shown with the answer.',
    'Omission of essential working will result in loss of marks.',
    'The use of an approved scientific calculator is expected, where appropriate.',
    'If the degree of accuracy is not specified in the question, and if the answer is not exact, '
    'give the answer to three significant figures. Give answers in degrees to one decimal place.',
    'For $\\pi$, use either your calculator value or 3.142, unless the question requires the answer in terms of $\\pi$.',
    'The number of marks is given in brackets [ ] at the end of each question or part question.',
]

FORMULAE_AM = [
    ('1.  ALGEBRA', None),
    ('Quadratic Equation', 'For the equation $ax^2 + bx + c = 0$, $x = \\dfrac{-b \\pm \\sqrt{b^2 - 4ac}}{2a}$'),
    ('Binomial expansion',
     '$(a + b)^n = a^n + \\binom{n}{1}a^{n-1}b + \\binom{n}{2}a^{n-2}b^2 + \\cdots + \\binom{n}{r}a^{n-r}b^r + \\cdots + b^n$, '
     'where $n$ is a positive integer and $\\binom{n}{r} = \\dfrac{n!}{r!(n-r)!} = \\dfrac{n(n-1)\\cdots(n-r+1)}{r!}$'),
    ('2.  TRIGONOMETRY', None),
    ('Identities', '$\\sin^2 A + \\cos^2 A = 1$'),
    ('', '$\\sec^2 A = 1 + \\tan^2 A$'),
    ('', '$\\operatorname{cosec}^2 A = 1 + \\cot^2 A$'),
    ('', '$\\sin(A \\pm B) = \\sin A \\cos B \\pm \\cos A \\sin B$'),
    ('', '$\\cos(A \\pm B) = \\cos A \\cos B \\mp \\sin A \\sin B$'),
    ('', '$\\tan(A \\pm B) = \\dfrac{\\tan A \\pm \\tan B}{1 \\mp \\tan A \\tan B}$'),
    ('', '$\\sin 2A = 2 \\sin A \\cos A$'),
    ('', '$\\cos 2A = \\cos^2 A - \\sin^2 A = 2\\cos^2 A - 1 = 1 - 2\\sin^2 A$'),
    ('', '$\\tan 2A = \\dfrac{2 \\tan A}{1 - \\tan^2 A}$'),
    ('Formulae for $\\triangle ABC$', '$\\dfrac{a}{\\sin A} = \\dfrac{b}{\\sin B} = \\dfrac{c}{\\sin C}$'),
    ('', '$a^2 = b^2 + c^2 - 2bc \\cos A$'),
    ('', '$\\Delta = \\dfrac{1}{2} bc \\sin A$'),
]

# 4052 (E Math): the formula list printed on the real paper.
FORMULAE_EM = [
    ('Compound interest', 'Total amount $= P\\left(1 + \\dfrac{r}{100}\\right)^{n}$'),
    ('Mensuration', 'Curved surface area of a cone $= \\pi r l$'),
    ('', 'Surface area of a sphere $= 4\\pi r^2$'),
    ('', 'Volume of a cone $= \\dfrac{1}{3}\\pi r^2 h$'),
    ('', 'Volume of a sphere $= \\dfrac{4}{3}\\pi r^3$'),
    ('', 'Area of triangle $ABC = \\dfrac{1}{2} ab \\sin C$'),
    ('', 'Arc length $= r\\theta$, where $\\theta$ is in radians'),
    ('', 'Sector area $= \\dfrac{1}{2} r^2 \\theta$, where $\\theta$ is in radians'),
    ('Trigonometry', '$\\dfrac{a}{\\sin A} = \\dfrac{b}{\\sin B} = \\dfrac{c}{\\sin C}$'),
    ('', '$a^2 = b^2 + c^2 - 2bc \\cos A$'),
    ('Statistics', 'Mean $= \\dfrac{\\sum fx}{\\sum f}$'),
    ('', 'Standard deviation $= \\sqrt{\\dfrac{\\sum fx^2}{\\sum f} - \\left(\\dfrac{\\sum fx}{\\sum f}\\right)^2}$'),
]


def formulae_for(shape):
    """The formula list the real paper of this syllabus prints (4049 → A Math, 4052 → E Math)."""
    return FORMULAE_EM if str(shape.get('code', '')) == '4052' else FORMULAE_AM


def size_math(doc):
    """Give every maths run the body's own size. A maths run with no size of its
    own is drawn at Word's 11 pt default by LibreOffice and by the iPhone's file
    preview, beside 9.5 pt text (Adrian, 21 Sep 2026: "the fonts are of different
    sizes?"). Runs that already carry a size (the grey solution notes) are left."""
    M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    half = str(int(round(doc.styles['Normal'].font.size.pt * 2)))
    for r in doc.element.body.iter(f'{{{M}}}r'):
        wrpr = r.find(qn('w:rPr'))
        if wrpr is None:
            wrpr = OxmlElement('w:rPr')
            mrpr = r.find(f'{{{M}}}rPr')
            if mrpr is not None:
                mrpr.addnext(wrpr)
            else:
                r.insert(0, wrpr)
        if wrpr.find(qn('w:sz')) is not None:
            continue
        if wrpr.find(qn('w:rFonts')) is None:
            rf = OxmlElement('w:rFonts')
            rf.set(qn('w:ascii'), 'Cambria Math')
            rf.set(qn('w:hAnsi'), 'Cambria Math')
            wrpr.insert(0, rf)
        for tag in ('w:sz', 'w:szCs'):
            e = OxmlElement(tag)
            e.set(qn('w:val'), half)
            wrpr.append(e)


def page_numbers(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    fld.append(r)
    p._element.append(fld)


def front_page(ws, paper, total):
    shape = paper.get('shape', {})
    code = f"{shape.get('code', '')}/0{paper.get('paperNo', '')}"
    ws.title(paper.get('title') or f"{shape.get('subject', 'Additional Mathematics').upper()}  {code}")
    ws.subtitle(f"{len(paper.get('questions') or [])} questions  ·  {total} marks  ·  {shape.get('duration', '')}")
    # A school-style paper (scripts/school-paper/run.mjs) carries its own front page:
    # front = {note, instructions[], formulae[]}; an empty formulae list prints no sheet.
    front = paper.get('front') or {}
    ws.para([('text', front.get('note') or 'Newly written questions in the GCE format, not a past-year paper.', {'italic': True})])
    ws.para([('text', '')])
    ws.para([('text', 'READ THESE INSTRUCTIONS FIRST', {'bold': True})])
    for line in front.get('instructions') or INSTRUCTIONS:
        ws.para(segs(line))
    ws.para(segs(f'The total number of marks for this paper is {total}.'))
    ws.page_break()
    if 'formulae' in front and not front['formulae']:
        return
    ws.para([('text', 'Mathematical Formulae', {'bold': True})])
    ws.para([('text', '')])
    for head, body in (front.get('formulae') or formulae_for(shape)):
        if body is None:
            p = ws.para([('text', head, {'bold': True})])
            p.paragraph_format.space_before = Cm(0.3)
            continue
        p = ws.para((segs(head, {'italic': True}) if head else []) + [('text', '\t')] + segs(body))
        p.paragraph_format.left_indent = Cm(4.0)
        p.paragraph_format.first_line_indent = Cm(-4.0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4.0))
    ws.page_break()


# ----------------------------------------------------------------- main ----

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('paper')
    ap.add_argument('--figures', default=None, help='directory holding Q<n>.png')
    ap.add_argument('--out', default=None, help='output directory (default: beside the JSON)')
    ap.add_argument('--space', type=float, default=3.0, help='writing lines per mark on the paper')
    a = ap.parse_args()

    paper = json.load(open(a.paper))
    slots = [s for s in paper['questions'] if s.get('accepted', True) and (s.get('question') or s.get('draft'))]
    total = sum(s['target'] for s in slots)
    name = splitext(basename(a.paper))[0]
    out_dir = a.out or dirname(abspath(a.paper))
    os.makedirs(out_dir, exist_ok=True)

    # --- the paper
    ws = Worksheet(working_space=a.space)
    layout = paper.get('layout') or {}
    shape = paper.get('shape', {})
    if layout.get('front_page') is False:
        # Adrian, 20 Sep 2026 (TJC): "just have title, then questions right away" —
        # no instructions page, no formulae sheet, no "[Answers for Question n]".
        ws.title(paper.get('title') or f"{shape.get('subject', 'Additional Mathematics').upper()}")
        ws.subtitle(f"{len(slots)} questions  ·  {total} marks  ·  {shape.get('duration', '')}")
    else:
        front_page(ws, paper, total)
    page_per_q = bool(layout.get('page_per_question'))
    for i, s in enumerate(slots):
        if page_per_q and i:
            ws.page_break()
        if page_per_q:
            # Every question on its own page. The blank space under each part is at least
            # `a.space` lines per mark (3 by default — Adrian, 20 Sep 2026: "you have to give
            # enough space"); the question takes as many whole pages as that needs (one, or
            # two for a long question) and the space is then stretched to FILL those pages
            # in proportion to marks, so no page is left nearly empty.
            q0 = s.get('question') or s.get('draft')
            used = estimate_used_cm(q0, a.figures, s['pos'], first_page=(i == 0))
            leaves = [x for p in (q0.get('parts') or []) for x in (p['subparts'] if p.get('subparts') else [p])]
            bonus = sum(1 for x in leaves if x.get('marks') == 1)      # one_mark_bonus lines
            min_lines = a.space * s['target'] + bonus
            pages = max(1, -(-(used + min_lines * LINE_CM + 0.8) // PAGE_USABLE_CM))
            avail = pages * PAGE_USABLE_CM - 0.8 - used
            ws.working_space = max(a.space, (avail / LINE_CM - bonus) / max(1, s['target']))
        q = question(ws, s, a.figures)
        got = sum((p.get('marks') or 0) if not p.get('subparts') else sum(x.get('marks') or 0 for x in p['subparts'])
                  for p in (q.get('parts') or [])) or s['target']
        if got != s['target']:
            print(f"  ⚠ Q{s['pos']} parts sum to {got}, slot target {s['target']}")
    page_numbers(ws.doc)
    size_math(ws.doc)
    paper_path = join(out_dir, name + '.docx')
    ws.save(paper_path)

    # --- worked solutions + answer key
    ws2 = Worksheet(working_space=0)
    shape = paper.get('shape', {})
    ws2.title(f"{paper.get('title') or shape.get('subject', 'Additional Mathematics').upper()}  ·  Worked solutions")
    ws2.subtitle(f"{total} marks")
    for s in slots:
        q = question(ws2, s, a.figures)
        ws2.solution_box(solution_rows(q.get('solution')), keep_together=False)
    ws2.section('Answers')
    for s in slots:
        q = s.get('question') or s.get('draft')
        p = ws2.para([('text', f"{s['pos']}\t", {'bold': True})] + segs(q.get('answer') or ''))
        p.paragraph_format.left_indent = Cm(Q_TEXT_CM)
        p.paragraph_format.first_line_indent = Cm(-Q_TEXT_CM)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(Q_TEXT_CM))
    page_numbers(ws2.doc)
    size_math(ws2.doc)
    sol_path = join(out_dir, name + '-solutions.docx')
    ws2.save(sol_path)

    if FAILED:
        print(f'  ⚠ {len(FAILED)} LaTeX fragment(s) pandoc could not convert were left as italic text:')
        for latex, err in FAILED[:10]:
            print('     ', latex[:90], '—', err)
    print(json.dumps({'paper': paper_path, 'solutions': sol_path, 'questions': len(slots), 'marks': total,
                      'math_fragments': len(_omml_cache), 'failed': len(FAILED)}))


if __name__ == '__main__':
    main()
